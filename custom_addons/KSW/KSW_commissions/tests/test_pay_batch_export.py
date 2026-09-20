"""Exporting a batch to Excel.

The case that asked for it is the imported one — sixty-seven Driver Trips
lines nobody can check against BAS inside a web form — so the fixture is
the BAS import's own, run and then exported. What is pinned here is that
the workbook opens, carries a row per entry, and shows the columns the
component actually uses (and none of the ones it does not).
"""
import base64
import io

import openpyxl

from .test_bas_trips_import import BasTripsImportCommon


class TestPayBatchExport(BasTripsImportCommon):

    def _export(self, batch):
        action = batch.action_export_xlsx()
        self.assertEqual(action['type'], 'ir.actions.act_url')
        att_id = int(action['url'].split('/web/content/')[1].split('?')[0])
        attachment = self.env['ir.attachment'].browse(att_id)
        self.assertTrue(attachment.name.endswith('.xlsx'))
        book = openpyxl.load_workbook(
            io.BytesIO(base64.b64decode(attachment.datas)))
        return attachment, book.active

    @staticmethod
    def _find_row(sheet, text):
        """Row index (1-based) whose first cell is ``text``."""
        for row in sheet.iter_rows(min_col=1, max_col=1):
            if row[0].value == text:
                return row[0].row
        return None

    def test_01_export_has_one_row_per_entry(self):
        batch = self._batch()
        self._import(batch)
        self.assertEqual(len(batch.entry_ids), 1)

        attachment, sheet = self._export(batch)

        header_row = self._find_row(sheet, 'Employee')
        self.assertTrue(header_row, 'the table header is missing')
        headers = [c.value for c in sheet[header_row]]
        # The driver's own row, then the totals row.
        self.assertEqual(sheet.cell(header_row + 1, 1).value,
                         self.driver_ok.name)
        self.assertEqual(sheet.cell(header_row + 2, 1).value, 'Total')

        entry = batch.entry_ids
        values = {headers[i]: sheet.cell(header_row + 1, i + 1).value
                  for i in range(len(headers))}
        self.assertAlmostEqual(values['Free Allowance'], entry.threshold_qty)
        self.assertAlmostEqual(values['Amount'], entry.amount, places=2)
        # The key the import matched on, so the figures can be traced back.
        self.assertEqual(values['BAS Cost Center'],
                         self.driver_ok.x_bas_driver_cost_center)

    def test_02_columns_follow_the_component(self):
        """A tiered, undated component exports no Date and no Type column."""
        batch = self._batch()
        self._import(batch)

        _attachment, sheet = self._export(batch)
        header_row = self._find_row(sheet, 'Employee')
        headers = [c.value for c in sheet[header_row]]

        self.assertNotIn('Date', headers)   # needs_date is False
        self.assertNotIn('Type', headers)   # has_options is False
        self.assertIn('Rate', headers)      # not a fixed-amount component

    def test_03_header_block_names_the_batch(self):
        batch = self._batch()
        self._import(batch)

        _attachment, sheet = self._export(batch)

        self.assertEqual(sheet.cell(self._find_row(sheet, 'Batch'), 2).value,
                         batch.name)
        self.assertEqual(
            sheet.cell(self._find_row(sheet, 'Component'), 2).value,
            batch.component_id.name)
        self.assertEqual(sheet.cell(self._find_row(sheet, 'Scope'), 2).value,
                         self.dept.name)

    def test_04_empty_batch_refuses(self):
        from odoo.exceptions import UserError
        batch = self._batch()
        with self.assertRaises(UserError):
            batch.action_export_xlsx()

    def test_05_export_survives_an_approved_batch(self):
        """Exporting is reading — it is not gated on the entry window."""
        batch = self._batch()
        self._import(batch)
        batch.sudo().write({'state': 'approved'})

        _attachment, sheet = self._export(batch)
        self.assertEqual(
            sheet.cell(self._find_row(sheet, 'Employee') + 1, 1).value,
            self.driver_ok.name)
