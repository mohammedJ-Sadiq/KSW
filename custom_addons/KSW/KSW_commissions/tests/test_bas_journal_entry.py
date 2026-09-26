"""The month as a BAS journal entry.

What is pinned here is the shape of the voucher the accountant used to
type by hand (قيد 26007264, July 2026): a driver's commission debited to
the commissions expense account with his truck in the cost-centre column,
his loan installment credited to his own BAS advance account, and only
the remainder credited to the accrual — and the whole thing balanced.

The failure that matters most is not a wrong number, it is a *missing*
one: the accrual line is derived as the residue, so a component with no
account configured would still produce a file that adds up while paying
the employee a figure that came out of nowhere. Those cases are tested as
refusals.
"""
import base64
import io

import openpyxl

from odoo.exceptions import UserError

from .test_commission_priority_deduction import CommissionPriorityCommon


class BasJournalCommon(CommissionPriorityCommon):

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        # Meals is the fixture's fixed-amount component and carries no
        # accounts of its own in the catalog; give it the overtime ones so
        # the happy path has something to post.
        cls.meals.sudo().write({
            'x_bas_expense_code': '3201010006',
            'x_bas_expense_name': 'مصروفات اضافي العاملين',
            'x_bas_accrual_code': '2107010001',
            'x_bas_accrual_name': 'مصروفات مستحقة عمولات سائقين التريلات',
        })
        cls.emp_a.sudo().x_loan_acc_no = '1205010385'
        cls.emp_b.sudo().x_loan_acc_no = '1205010386'
        # A second employee in the *same* department: the fixture's
        # supervisor may only record pay for his own (which is the point of
        # the scoping), so a two-employee entry needs him here.
        cls.emp_a2 = cls._employee('Sub Emp A2', cls.dept_a, 6000.0)
        cls.emp_a2.sudo().x_loan_acc_no = '1205010387'

    def _rows(self, run):
        return run._bas_journal_rows()

    def _voucher_name(self, component):
        """The component's name as the voucher writes it.

        The export builds in the voucher's language, not the session's
        (``ksw.bas.journal.voucher_lang``), so a test that compares against
        ``component.name`` is comparing against the wrong string the moment
        an Arabic translation exists.
        """
        lang = self.env['ksw.bas.journal'].voucher_lang()
        return component.with_context(lang=lang).name

    @staticmethod
    def _by_code(rows, code):
        return [r for r in rows if r['code'] == code]


class TestJournalRows(BasJournalCommon):

    def test_01_one_debit_and_one_credit_for_a_plain_month(self):
        batch = self._commission_entry(self.emp_a, 1200.0).batch_id
        run = self._approve(batch)

        rows = self._rows(run)

        self.assertEqual(len(rows), 2)
        debit, credit = rows
        self.assertEqual(debit['code'], '3201010006')
        self.assertAlmostEqual(debit['debit'], 1200.0)
        self.assertFalse(debit['credit'])
        self.assertEqual(credit['code'], '2107010001')
        self.assertAlmostEqual(credit['credit'], 1200.0)
        self.assertIn(self.emp_a.name, debit['ref'])

    def test_02_loan_comes_out_before_the_accrual(self):
        """The three-line block the hand-typed voucher used."""
        self._make_pending_installment(self.emp_a, 500.0)
        batch = self._commission_entry(self.emp_a, 1200.0).batch_id
        run = self._approve(batch)

        rows = self._rows(run)

        self.assertEqual([r['code'] for r in rows],
                         ['3201010006', '1205010385', '2107010001'])
        self.assertAlmostEqual(rows[0]['debit'], 1200.0)
        self.assertAlmostEqual(rows[1]['credit'], 500.0)
        # Only what is left is owed to him.
        self.assertAlmostEqual(rows[2]['credit'], 700.0)

    def test_03_a_loan_that_eats_the_whole_month_leaves_no_accrual(self):
        self._make_pending_installment(self.emp_a, 1200.0)
        batch = self._commission_entry(self.emp_a, 1200.0).batch_id
        run = self._approve(batch)

        rows = self._rows(run)

        self.assertEqual([r['code'] for r in rows],
                         ['3201010006', '1205010385'])
        self.assertAlmostEqual(rows[1]['credit'], 1200.0)

    def test_04_the_entry_balances(self):
        self._make_pending_installment(self.emp_a, 500.0)
        # Both on one batch: a scope+component+month has exactly one.
        batch = self._batch(self.sup_a, self.dept_a, component=self.meals)
        self._entry(batch, self.emp_a, user=self.sup_a,
                    quantity=1.0, amount_override=1200.0)
        self._entry(batch, self.emp_a2, user=self.sup_a,
                    quantity=1.0, amount_override=800.0)
        run = self._approve(batch)

        rows = self._rows(run)

        self.assertAlmostEqual(sum(r['debit'] for r in rows),
                               sum(r['credit'] for r in rows), places=2)
        self.env['ksw.bas.journal'].check_balanced(rows)

    def test_05_cost_centre_is_the_truck_where_the_component_asks(self):
        # The vehicle (BAS COST_CENTER, filled by the pull), never the
        # driver's own cost centre — that one identifies the man.
        self.emp_a.sudo().x_bas_equipment_code = 'T166'
        self.emp_a.sudo().x_bas_driver_cost_center = 'AYAZ KHAN DURANI385'
        batch = self._commission_entry(self.emp_a, 1200.0).batch_id
        run = self._approve(batch)

        # Meals does not stamp it...
        self.assertEqual(self._rows(run)[0]['cost_center'], '')

        # ...the same entry on a component that does, carries the truck.
        self.meals.sudo().x_bas_use_cost_center = True
        self.assertEqual(self._rows(run)[0]['cost_center'], 'T166')


class TestJournalDescriptions(BasJournalCommon):
    """What the line *says* — the half of the voucher nobody can compute."""

    def test_30_the_supervisors_note_is_the_description(self):
        batch = self._batch(self.sup_a, self.dept_a, component=self.meals)
        self._entry(batch, self.emp_a, user=self.sup_a, quantity=1.0,
                    amount_override=300.0, reason='تصليح قفل باب سيارة ١٢٤')
        run = self._approve(batch)

        debit = self._rows(run)[0]

        # His words, not the component's name.
        self.assertIn('تصليح قفل باب سيارة ١٢٤', debit['ref'])
        self.assertNotIn(self._voucher_name(self.meals), debit['ref'])

    def test_31_a_line_per_note_not_per_component(self):
        """Three notes, three lines — one accrual credit for the lot."""
        batch = self._batch(self.sup_a, self.dept_a, component=self.meals)
        for note, amount in (('قير 108', 100.0), ('سير 158', 200.0),
                             ('ماطور 131', 300.0)):
            self._entry(batch, self.emp_a, user=self.sup_a, quantity=1.0,
                        amount_override=amount, reason=note)
        run = self._approve(batch)

        rows = self._rows(run)

        self.assertEqual(len(rows), 4)
        self.assertEqual([r['debit'] for r in rows[:3]], [100.0, 200.0, 300.0])
        for note, row in zip(('قير 108', 'سير 158', 'ماطور 131'), rows[:3]):
            self.assertIn(note, row['ref'])
        self.assertAlmostEqual(rows[3]['credit'], 600.0)

    def test_32_occurrences_described_the_same_way_stay_one_line(self):
        batch = self._batch(self.sup_a, self.dept_a, component=self.meals)
        for _i in range(3):
            self._entry(batch, self.emp_a, user=self.sup_a, quantity=1.0,
                        amount_override=100.0, reason='قير 108')
        run = self._approve(batch)

        rows = self._rows(run)

        self.assertEqual(len(rows), 2)
        self.assertAlmostEqual(rows[0]['debit'], 300.0)

    def test_33_no_note_falls_back_to_the_component_name(self):
        batch = self._batch(self.sup_a, self.dept_a, component=self.meals)
        self._entry(batch, self.emp_a, user=self.sup_a, quantity=1.0,
                    amount_override=300.0, reason='')
        run = self._approve(batch)

        self.assertIn(self._voucher_name(self.meals),
                      self._rows(run)[0]['ref'])

    def test_33b_an_importers_audit_note_is_not_a_description(self):
        """`details` on an imported component is the machine's, not his."""
        self.meals.sudo().importer = 'bas_trips'
        batch = self._batch(self.sup_a, self.dept_a, component=self.meals)
        self._entry(batch, self.emp_a, user=self.sup_a, quantity=1.0,
                    amount_override=300.0, reason='',
                    details='Weighted on «الرد المضاعف». Worked days: 31.')
        run = self._approve(batch)

        ref = self._rows(run)[0]['ref']

        self.assertNotIn('Worked days', ref)
        self.assertIn(self._voucher_name(self.meals), ref)

    def test_33c_the_voucher_is_arabic_whoever_exports_it(self):
        """An English session must still produce an Arabic voucher."""
        arabic = self.env['res.lang'].sudo().search(
            [('code', '=like', 'ar%'), ('active', '=', True)], limit=1)
        if not arabic:
            self.skipTest('no Arabic language installed on this database')
        self.meals.with_context(lang=arabic.code).sudo().name = 'الوجبات'
        batch = self._batch(self.sup_a, self.dept_a, component=self.meals)
        self._entry(batch, self.emp_a, user=self.sup_a, quantity=1.0,
                    amount_override=300.0, reason='')
        run = self._approve(batch)

        rows = run.with_context(lang='en_US')._bas_journal_rows()

        self.assertIn('الوجبات', rows[0]['ref'])

    def test_34_the_name_is_the_one_bas_holds(self):
        """BAS's own spelling, minus the employee number it appends."""
        self.emp_a.sudo().x_bas_driver_cost_center = 'عبدالله محمد عبد الحفيظ محمد700'
        batch = self._commission_entry(self.emp_a, 500.0).batch_id
        run = self._approve(batch)

        ref = self._rows(run)[0]['ref']

        self.assertIn('عبدالله محمد عبد الحفيظ محمد', ref)
        self.assertNotIn('700', ref)
        self.assertNotIn(self.emp_a.name, ref)

    def test_35_without_a_bas_name_his_odoo_name_is_used(self):
        self.emp_a.sudo().x_bas_driver_cost_center = False
        batch = self._commission_entry(self.emp_a, 500.0).batch_id
        run = self._approve(batch)

        self.assertIn(self.emp_a.name, self._rows(run)[0]['ref'])


class TestJournalRefusals(BasJournalCommon):

    def test_10_component_without_an_account_is_named(self):
        self.meals.sudo().x_bas_expense_code = False
        batch = self._commission_entry(self.emp_a, 1200.0).batch_id
        run = self._approve(batch)

        with self.assertRaises(UserError) as caught:
            self._rows(run)
        # Named in the voucher's language, like every other word on it.
        self.assertIn(self._voucher_name(self.meals), str(caught.exception))

    def test_11_employee_without_a_loan_account_is_named(self):
        self.emp_a.sudo().x_loan_acc_no = False
        self._make_pending_installment(self.emp_a, 500.0)
        batch = self._commission_entry(self.emp_a, 1200.0).batch_id
        run = self._approve(batch)

        with self.assertRaises(UserError) as caught:
            self._rows(run)
        self.assertIn(self.emp_a.name, str(caught.exception))

    def test_12_a_register_figure_its_entries_do_not_explain(self):
        """A hand-corrected register line has no accounts to post to."""
        batch = self._commission_entry(self.emp_a, 1200.0).batch_id
        run = self._approve(batch)
        run.line_ids.sudo().write({'earnings': 5000.0})

        with self.assertRaises(UserError) as caught:
            self._rows(run)
        self.assertIn(self.emp_a.name, str(caught.exception))


class TestJournalWorkbook(BasJournalCommon):

    def _sheet(self, run):
        data = run._bas_journal_workbook()
        return openpyxl.load_workbook(io.BytesIO(data)).active

    def test_20_the_layout_is_the_one_bas_exports(self):
        batch = self._commission_entry(self.emp_a, 1200.0).batch_id
        run = self._approve(batch)

        sheet = self._sheet(run)

        # Row 1 empty, the nine keys on row 2, data from row 3.
        self.assertTrue(all(c.value is None for c in sheet[1]))
        self.assertEqual(
            [c.value for c in sheet[2]],
            ['number', 'fdate', 'fcode', 'fname', 'amount', 'tamount',
             'ref', 'remark', 'cost_Center'])
        self.assertTrue(sheet.sheet_view.rightToLeft)

        # The voucher number is BAS's to assign, so the column is blank.
        self.assertEqual(sheet.cell(3, 1).value.strip(), '')
        # Period end, not the day it was exported.
        self.assertEqual(sheet.cell(3, 2).value.strip(), '31/03/2029')
        # Codes are numbers, text is padded, an empty cell is two spaces.
        self.assertEqual(sheet.cell(3, 3).value, 3201010006)
        self.assertEqual(sheet.cell(3, 4).value,
                         ' مصروفات اضافي العاملين ')
        self.assertEqual(sheet.cell(3, 5).value, 1200.0)
        self.assertEqual(sheet.cell(3, 6).value, '  ')
        self.assertEqual(sheet.cell(3, 8).value, '  ')

    def test_21_the_export_wizard_offers_it(self):
        batch = self._commission_entry(self.emp_a, 1200.0).batch_id
        run = self._approve(batch)

        wizard = self.env['ksw.commission.bank.export.wizard'].sudo().create({
            'run_id': run.id, 'export_mode': 'journal_entry'})
        action = wizard.action_export()

        attachment = self.env['ir.attachment'].browse(
            int(action['url'].split('/web/content/')[1].split('?')[0]))
        self.assertTrue(attachment.name.startswith('JournalEntry_'))
        sheet = openpyxl.load_workbook(
            io.BytesIO(base64.b64decode(attachment.datas))).active
        self.assertEqual(sheet.cell(2, 1).value, 'number')
