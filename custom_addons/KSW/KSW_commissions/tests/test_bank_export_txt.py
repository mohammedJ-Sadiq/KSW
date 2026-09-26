"""The Kawthar TXT commission export — it had never produced a file.

`_make_kawthar_txt(self, bank, lines)` opened with `lines = []`, rebinding its
own parameter to an empty list, so the very next statement — `lines.sorted(…)`
— raised `'list' object has no attribute 'sorted'`. Every call died on the
first line of the loop, from both `_export_all_txt` and `_export_specific_txt`.

Nothing caught it because this wizard had no test at all: the module's export
coverage was the *batch* xlsx (`test_pay_batch_export.py`), not the bank file.
So the whole body below the crash had never executed either — which is why
this test asserts the 194-character row contract and the field reads, not just
"it does not raise".

Hit on KSWCO 2026-09-24 by the accountant, on the August run.
"""
from datetime import date

from odoo.tests.common import TransactionCase


class TestKawtharTxtExport(TransactionCase):

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        env = cls.env
        Bank = env['res.partner.bank']
        if 'x_file_type' not in Bank._fields:
            cls.skip_reason = 'KSW_payroll not installed — no x_file_type'
            return
        cls.skip_reason = None

        cls.dept = env['hr.department'].sudo().create({'name': 'TXT Dept'})
        cls.partner = env['res.partner'].sudo().create({'name': 'TXT Bank Co'})
        cls.bank = Bank.sudo().create({
            'acc_number': 'SA0380000000608010167519',
            'partner_id': cls.partner.id,
            'x_file_type': 'kawthar',
            'x_wps_cic_number': '1234567890',
        })
        cls.emp = env['hr.employee'].sudo().create({
            'name': 'TXT Employee',
            'department_id': cls.dept.id,
            'barcode': '100000000042',
        })
        # The employee's own Kawthar card: CIC(5) + card number(14).
        cls.card = Bank.sudo().create({
            'acc_number': '1234512345678901234',
            'partner_id': cls.emp.work_contact_id.id,
        })
        cls.emp.sudo().write({
            'x_salary_bank_account_id': cls.bank.id,
            'bank_account_ids': [(4, cls.card.id)],
            'x_employee_no': '4242',
            'ssnid': '1098765432',
        })

        cls.pay_run = env['ksw.pay.run'].sudo().create({'period': '2029-05-01'})
        cls.line = env['ksw.pay.run.line'].sudo().create({
            'run_id': cls.pay_run.id,
            'employee_id': cls.emp.id,
            'earnings': 1500.0,
            'loan_offset': 250.0,
        })

    def _wizard(self, mode):
        return self.env['ksw.commission.bank.export.wizard'].sudo().create({
            'run_id': self.pay_run.id,
            'export_mode': mode,
            'bank_account_id': self.bank.id,
            'value_date': date(2029, 5, 28),
            'operation_code': '2',
        })

    # ------------------------------------------------------------------
    def test_the_txt_body_is_produced_at_all(self):
        """The regression itself: this used to raise AttributeError."""
        if self.skip_reason:
            self.skipTest(self.skip_reason)
        text = self._wizard('specific_txt')._make_kawthar_txt(
            self.bank, self.pay_run.line_ids)
        self.assertTrue(text, 'the export produced an empty file')

    def test_every_row_is_exactly_194_characters(self):
        if self.skip_reason:
            self.skipTest(self.skip_reason)
        text = self._wizard('specific_txt')._make_kawthar_txt(
            self.bank, self.pay_run.line_ids)
        rows = text.split('\n')
        self.assertEqual(len(rows), 1, 'one payable employee, one row')
        for row in rows:
            self.assertEqual(len(row), 194)

    def test_the_row_carries_the_employee_and_the_amount(self):
        if self.skip_reason:
            self.skipTest(self.skip_reason)
        text = self._wizard('specific_txt')._make_kawthar_txt(
            self.bank, self.pay_run.line_ids)
        row = text.split('\n')[0]
        self.assertEqual(row[:12], '000000000001', 'running 1..N, as payroll')
        self.assertEqual(row[12:22], '0000012345', 'CIC from the card')
        self.assertEqual(row[22:36], '12345678901234', 'the card number')
        self.assertEqual(row[86:96], '1098765432', 'the SSN (ssnid)')
        self.assertIn('TXT Employee', text, 'the employee name')
        # net_payable = earnings - loan_offset = 1250.00 → 125000 halalas
        self.assertIn(str(125000).zfill(15), text, 'the net in halalas')
        self.assertIn('20290528', text, 'the value date')

    def test_a_zero_or_negative_line_is_left_out(self):
        if self.skip_reason:
            self.skipTest(self.skip_reason)
        self.env['ksw.pay.run.line'].sudo().create({
            'run_id': self.pay_run.id,
            'employee_id': self.env['hr.employee'].sudo().create({
                'name': 'TXT Zero', 'department_id': self.dept.id,
                'barcode': '100000000043',
                'bank_account_ids': [(4, self.card.id)]}).id,
            'earnings': 0.0,
        })
        text = self._wizard('specific_txt')._make_kawthar_txt(
            self.bank, self.pay_run.line_ids)
        self.assertNotIn('TXT Zero', text)
        self.assertEqual(len(text.split('\n')), 1)

    def test_export_all_txt_runs_end_to_end(self):
        """The path the accountant actually pressed."""
        if self.skip_reason:
            self.skipTest(self.skip_reason)
        action = self._wizard('all_txt').action_export()
        self.assertTrue(action, 'the wizard returned no download action')
        att = self.env['ir.attachment'].sudo().search(
            [('res_model', '=', 'ksw.pay.run'), ('res_id', '=', self.pay_run.id)],
            order='id desc', limit=1)
        self.assertTrue(att, 'no file was attached')
        self.assertTrue(att.name.endswith('.txt'), att.name)

    def test_a_line_without_an_employee_account_is_left_out(self):
        """The company account the transfer leaves from is not a card."""
        if self.skip_reason:
            self.skipTest(self.skip_reason)
        self.emp.sudo().write({'bank_account_ids': [(5,)]})
        text = self._wizard('specific_txt')._make_kawthar_txt(
            self.bank, self.pay_run.line_ids)
        self.assertNotIn('SA0380000000608010167519'[5:19], text)
        self.assertFalse(text)


class TestCommissionBankExcel(TestKawtharTxtExport):
    """Employee No / SSN on both sheets, and the merged workbook."""

    def _sheets(self, data):
        import io
        import openpyxl
        return openpyxl.load_workbook(io.BytesIO(data))

    def test_excel_carries_employee_number_and_ssn(self):
        if self.skip_reason:
            self.skipTest(self.skip_reason)
        wiz = self._wizard('specific_excel')
        wb = self._sheets(wiz._make_wps_excel({self.bank: self.pay_run.line_ids}))
        summary = [c.value for c in wb['Commission Summary'][2]]
        self.assertIn('4242', summary)
        self.assertIn('1098765432', summary)
        wps = [c.value for c in wb['WPS'][7]]
        self.assertEqual(wps[1], '1234512345678901234', "the employee's card")
        self.assertEqual(wps[3], '4242')
        self.assertEqual(wps[4], '1098765432')

    def test_all_excel_merges_banks_into_one_file(self):
        if self.skip_reason:
            self.skipTest(self.skip_reason)
        bank2 = self.env['res.partner.bank'].sudo().create({
            'acc_number': 'SA0000000000000000000002',
            'partner_id': self.partner.id,
            'x_file_type': 'wps',
        })
        emp2 = self.env['hr.employee'].sudo().create({
            'name': 'XLS Other', 'department_id': self.dept.id})
        self.env['ksw.pay.run.line'].sudo().create({
            'run_id': self.pay_run.id, 'employee_id': emp2.id,
            'earnings': 900.0, 'bank_account_id': bank2.id,
        })
        self._wizard('all_excel').action_export()
        att = self.env['ir.attachment'].sudo().search(
            [('res_model', '=', 'ksw.pay.run'), ('res_id', '=', self.pay_run.id)],
            order='id desc', limit=1)
        self.assertTrue(att.name.endswith('.xlsx'), att.name)
        wb = self._sheets(att.raw)
        self.assertEqual(len(wb.sheetnames), 3, wb.sheetnames)
        self.assertEqual(wb['Commission Summary'].max_row, 3)

    def test_all_excel_split_still_zips_one_file_per_bank(self):
        if self.skip_reason:
            self.skipTest(self.skip_reason)
        bank2 = self.env['res.partner.bank'].sudo().create({
            'acc_number': 'SA0000000000000000000003',
            'partner_id': self.partner.id,
            'x_file_type': 'wps',
        })
        emp2 = self.env['hr.employee'].sudo().create({
            'name': 'XLS Split', 'department_id': self.dept.id})
        self.env['ksw.pay.run.line'].sudo().create({
            'run_id': self.pay_run.id, 'employee_id': emp2.id,
            'earnings': 900.0, 'bank_account_id': bank2.id,
        })
        wiz = self._wizard('all_excel')
        wiz.excel_layout = 'split'
        wiz.action_export()
        att = self.env['ir.attachment'].sudo().search(
            [('res_model', '=', 'ksw.pay.run'), ('res_id', '=', self.pay_run.id)],
            order='id desc', limit=1)
        self.assertTrue(att.name.endswith('.zip'), att.name)
