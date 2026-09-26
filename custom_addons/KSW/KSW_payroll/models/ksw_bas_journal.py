"""The journal entry, in the shape BAS reads.

BAS9 takes a posting as a flat spreadsheet — one row per journal line, nine
columns, debits and credits in two separate columns — and that is the only
interface it offers.  Until now the accountant retyped a whole month of
commissions and payroll into it by hand, one line at a time.

This model is the writer, shared by everything that has to produce such a
file: the monthly commission run, the payslip batch, anything added later.
It owns three things and deliberately nothing else:

``row()``
    one journal line, as a dict.
``employee_block()``
    the pattern every KSW posting follows — a person's earnings debited to
    their expense accounts, what he owes credited to his loan account, the
    rest credited to the accrual/payable account his earnings belong to.
``build_workbook()``
    the bytes, in the exact layout BAS produced: an empty first row, the
    nine header keys on row 2, right-to-left, text cells padded with a
    space at each end and empty ones written as two spaces.

That layout is copied from a real BAS export rather than guessed, because
an import template that "looks the same" is the one class of bug nobody
finds until the accountant is standing in front of the closing.

The *accounts* are not here.  They belong to whatever is being posted — a
pay component knows which expense account its money comes out of, a salary
rule knows its own — and this model never guesses one.  A figure with no
account configured is refused by name, never silently dropped: a journal
entry that balances because a line went missing is worse than no file.
"""
import io
import re

from odoo import _, api, models
from odoo.exceptions import UserError

try:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
except ImportError:  # pragma: no cover - openpyxl ships with Odoo
    openpyxl = None

try:
    from babel.dates import format_date as _babel_format_date
except ImportError:  # pragma: no cover
    _babel_format_date = None

# The nine columns, in BAS's own order and spelling. ``cost_Center`` really
# is capitalised in the middle — matched exactly, not tidied.
COLUMNS = [
    'number', 'fdate', 'fcode', 'fname', 'amount', 'tamount', 'ref',
    'remark', 'cost_Center',
]

# BAS writes an empty text cell as two spaces and pads every other one with
# a single space at each end. Mirrored so the file we hand back is
# byte-for-byte the kind of thing its importer already accepts.
EMPTY = '  '

# BAS writes the employee's number on the end of his name in the cost
# centre ("عبدالله محمد عبد الحفيظ محمد700"). The voucher carries the name
# alone, so the trailing number comes off.
_TRAILING_NUMBER = re.compile(r'[\s\-_]*\d+\s*$')

# Everything is rounded to two places before it is written, and the
# balancing line is derived from the rounded figures rather than recomputed
# — see ``employee_block``.
ROUND = 2


class KswBasJournal(models.AbstractModel):
    _name = 'ksw.bas.journal'
    _description = 'BAS Journal Entry Export'

    # ------------------------------------------------------------------
    # One line
    # ------------------------------------------------------------------
    @api.model
    def row(self, code, name, debit=0.0, credit=0.0, ref='', cost_center=''):
        """One journal line.

        ``code``/``name`` are the BAS account; exactly one of ``debit`` /
        ``credit`` carries the figure.
        """
        return {
            'code': code or '',
            'name': name or '',
            'debit': round(debit or 0.0, ROUND),
            'credit': round(credit or 0.0, ROUND),
            'ref': ref or '',
            'cost_center': cost_center or '',
        }

    # ------------------------------------------------------------------
    # One person's posting
    # ------------------------------------------------------------------
    @api.model
    def employee_block(self, debits, credits=None):
        """The rows for one employee: what he earned, what he owed, the rest.

        ``debits`` is a list of dicts, one per thing he is being paid for::

            {'code', 'name', 'amount', 'ref', 'cost_center',
             'credit_code', 'credit_name', 'credit_ref'}

        The ``credit_*`` keys say which accrual (commissions) or payable
        (payroll) account that money becomes a liability on once expensed.
        Entries sharing an account are credited as one line, in the order
        they first appear — which is what the hand-typed entry did.

        ``credits`` is what is taken out of the payment before it is owed
        to him — settled loan installments, GOSI, a penalty — each
        ``{'code', 'name', 'amount', 'ref'}``. They are consumed against
        the accrual groups in order, so an employee paid out of two
        accounts has his deduction taken off the first until it runs out.
        Whatever is left of each group is credited to it.

        A person who owes more than he earned this month leaves a negative
        residue. It is written as a *debit* on the payable account rather
        than a negative credit: BAS has no negative-credit concept, and the
        entry still balances.
        """
        rows = []
        groups = []          # [(credit_code, credit_name, credit_ref, total)]
        index = {}           # credit_code -> position in ``groups``
        for debit in debits:
            amount = round(debit.get('amount') or 0.0, ROUND)
            if not amount:
                continue
            rows.append(self.row(
                debit['code'], debit.get('name'), debit=amount,
                ref=debit.get('ref'), cost_center=debit.get('cost_center'),
            ))
            code = debit.get('credit_code') or ''
            if code not in index:
                index[code] = len(groups)
                groups.append([code, debit.get('credit_name') or '',
                               debit.get('credit_ref') or '', 0.0])
            groups[index[code]][3] = round(groups[index[code]][3] + amount,
                                           ROUND)

        for credit in (credits or []):
            amount = round(credit.get('amount') or 0.0, ROUND)
            if not amount:
                continue
            rows.append(self.row(
                credit['code'], credit.get('name'), credit=amount,
                ref=credit.get('ref'),
            ))
            # Take it off the accrual groups, first one first.
            for group in groups:
                if amount <= 0:
                    break
                # max(…, 0): a previous credit may already have pushed
                # the last group negative, and taking a negative amount
                # off it would give money back.
                take = min(max(group[3], 0.0), amount)
                group[3] = round(group[3] - take, ROUND)
                amount = round(amount - take, ROUND)
            if amount > 0 and groups:
                # More owed than earned: the last group goes negative and
                # the residue below turns into a debit.
                groups[-1][3] = round(groups[-1][3] - amount, ROUND)

        for code, name, ref, residue in groups:
            if not residue:
                continue
            if residue > 0:
                rows.append(self.row(code, name, credit=residue, ref=ref))
            else:
                rows.append(self.row(code, name, debit=-residue, ref=ref))
        return rows

    # ------------------------------------------------------------------
    # Checks
    # ------------------------------------------------------------------
    @api.model
    def check_balanced(self, rows):
        """Refuse an entry that does not balance, rather than export it."""
        debit = round(sum(r['debit'] for r in rows), ROUND)
        credit = round(sum(r['credit'] for r in rows), ROUND)
        if abs(debit - credit) >= 0.005:
            raise UserError(_(
                'The journal entry does not balance: debits %(debit).2f '
                'against credits %(credit).2f. Nothing was exported — this '
                'is a bug, please report it with the period.',
                debit=debit, credit=credit))
        return True

    @api.model
    def check_accounts(self, rows):
        """Every line needs an account. Refuse by name, never drop."""
        missing = [r for r in rows if not r['code']]
        if missing:
            raise UserError(_(
                'These amounts have no BAS account configured, so the '
                'entry cannot be written:\n%(what)s',
                what='\n'.join(
                    '• %s — %.2f' % (r['ref'] or r['name'] or '?',
                                     r['debit'] or r['credit'])
                    for r in missing[:20])))
        return True

    # ------------------------------------------------------------------
    # Accounts
    # ------------------------------------------------------------------
    @api.model
    def voucher_lang(self):
        """The language the voucher is written in — Arabic, if installed.

        Not the exporting user's. This file is read in BAS, next to BAS's
        own Arabic records, by an accountant who does not care which
        language the person who pressed the button happens to use. An
        export whose wording depends on that is an export that produces a
        different voucher for two people on the same month.

        Falls back to whatever is in context when no Arabic language is
        installed, so nothing breaks on an English-only database.
        """
        lang = self.env['res.lang'].sudo().search(
            [('code', '=like', 'ar%'), ('active', '=', True)],
            order='code', limit=1)
        return lang.code or self.env.context.get('lang')

    @api.model
    def employee_label(self, employee):
        """The employee as BAS knows him.

        The voucher is read in BAS, next to BAS's own records, so the name
        on it has to be the one BAS holds — that is
        ``x_bas_driver_cost_center``, the field kept for exactly this
        («مركز تكلفة الموظف»), minus the employee number BAS appends to
        it. Where it is not filled in, his Odoo name is all there is.

        ``getattr`` because the field belongs to KSW_commissions, which
        depends on this module and therefore loads after it — the same
        reason ``loan_account`` reads ``x_loan_acc_no`` that way.
        """
        bas_name = (getattr(employee, 'x_bas_driver_cost_center', '')
                    or '').strip()
        if bas_name:
            stripped = _TRAILING_NUMBER.sub('', bas_name).strip()
            if stripped:
                return stripped
        return employee.name or ''

    @api.model
    def check_loan_accounts(self, employees):
        """Every employee whose repayment has nowhere to go, in one list.

        Named together rather than one export at a time: the accountant
        fixes the employee records in a single visit, the way the BAS
        importer reports its skipped drivers.
        """
        missing = [employee for employee in employees
                   if not (getattr(employee, 'x_loan_acc_no', '') or '').strip()]
        if missing:
            raise UserError(_(
                'These employees have an installment settled in this period '
                'but no "Loan Acc No. in Bas" on their employee record, so '
                'the repayment has no account to be credited to:\n\n'
                '%(who)s\n\nSet it on each of them and export again.',
                who='\n'.join('• %s' % (e.name or '') for e in missing)))
        return True

    @api.model
    def loan_account(self, employee):
        """The employee's own loan account in BAS, and its name there.

        The code is always the employee's, never a rule's or a component's:
        one man's installment cannot be credited to another's account. The
        name comes from the synced BAS chart when KSW_ext_sync is installed
        — it mirrors the 120501 advances exactly — and is made up from his
        own name otherwise.

        ``employee`` must already be sudo'd: ``x_loan_acc_no`` is gated to
        ``hr.group_hr_user`` and the accountant exporting this is not an HR
        user.
        """
        # getattr, not a plain read: ``x_loan_acc_no`` is declared by
        # KSW_deduction, which depends on this module and therefore loads
        # *after* it. In a running database it is always there; during this
        # module's own load (its tests, a migration) it is not, and the
        # missing-account message below is the right answer either way.
        code = (getattr(employee, 'x_loan_acc_no', '') or '').strip()
        if not code:
            raise UserError(_(
                '%(name)s has an installment settled in this period but no '
                '"Loan Acc No. in Bas" on his employee record, so the '
                'repayment has no account to be credited to.',
                name=employee.name or ''))
        name = ''
        if 'ksw.bas.account' in self.env:
            account = self.env['ksw.bas.account'].sudo().search(
                [('bas_code', '=', code)], limit=1)
            name = account.name_ar or account.name_en or ''
        return code, name or _('Advance — %(name)s', name=employee.name or '')

    # ------------------------------------------------------------------
    # Labels
    # ------------------------------------------------------------------
    @api.model
    def ref(self, *parts):
        """A description line: the parts that have something in them.

        Joined on single spaces with each part stripped — employee names
        and rule labels in this database carry stray spacing, and a
        voucher line reading "Basic Salary  AHMED  June" is the sort of
        thing the accountant has to tidy by hand after every import.
        """
        return ' '.join(str(part).strip() for part in parts
                        if part and str(part).strip())
    # ------------------------------------------------------------------
    # Labels
    # ------------------------------------------------------------------
    @api.model
    def month_label(self, date):
        """The month as the accountant writes it — ``يوليو``, ``July``.

        No year: the hand-written entry never carried one, and the date
        column already does.
        """
        if not date:
            return ''
        lang = self.env.context.get('lang') or self.env.user.lang or 'en_US'
        if _babel_format_date:
            try:
                return _babel_format_date(date, 'MMMM', locale=lang)
            except Exception:
                pass
        return date.strftime('%B')

    # ------------------------------------------------------------------
    # The file
    # ------------------------------------------------------------------
    @api.model
    def build_workbook(self, rows, date, sheet_name=None, number=''):
        """The nine-column sheet BAS imports.

        ``number`` is left empty by design: BAS assigns the voucher number
        when the entry is imported.
        """
        if not openpyxl:
            raise UserError(_('openpyxl is required for the Excel export.'))
        self.check_accounts(rows)
        self.check_balanced(rows)

        book = openpyxl.Workbook()
        sheet = book.active
        sheet.title = (sheet_name or 'Journal')[:31]
        sheet.sheet_view.rightToLeft = True

        head = Font(bold=True)
        fill = PatternFill('solid', fgColor='D9E1F2')
        for column, key in enumerate(COLUMNS, 1):
            cell = sheet.cell(2, column, key)
            cell.font = head
            cell.fill = fill
            cell.alignment = Alignment(horizontal='center')

        fdate = ' %s ' % date.strftime('%d/%m/%Y') if date else EMPTY
        line = 3
        for row in rows:
            code = row['code']
            sheet.cell(line, 1, number or EMPTY)
            sheet.cell(line, 2, fdate)
            sheet.cell(line, 3, int(code) if str(code).isdigit() else code)
            sheet.cell(line, 4, self._pad(row['name']))
            sheet.cell(line, 5, row['debit'] or EMPTY)
            sheet.cell(line, 6, row['credit'] or EMPTY)
            sheet.cell(line, 7, self._pad(row['ref']))
            sheet.cell(line, 8, EMPTY)
            sheet.cell(line, 9, self._pad(row['cost_center']))
            line += 1

        for column, width in zip('ABCDEFGHI',
                                 (12, 14, 14, 42, 12, 12, 46, 10, 14)):
            sheet.column_dimensions[column].width = width

        buffer = io.BytesIO()
        book.save(buffer)
        return buffer.getvalue()

    @api.model
    def _pad(self, value):
        """A space at each end, as BAS's own export writes it."""
        value = (value or '').strip()
        return ' %s ' % value if value else EMPTY
