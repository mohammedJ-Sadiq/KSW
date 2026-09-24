"""The payslip batch as a BAS journal entry.

Same voucher shape as the commission run (see KSW_commissions'
``test_bas_journal_entry``), from a different source: the salary rules say
which account each figure lands on, and the employee's net is the residue.

The test that earns its place is :meth:`test_11_an_unmapped_rule_refuses`.
A rule left unposted does *not* unbalance the file — the payable line is
derived from what was posted, so the entry would still add up while
crediting the employee a net his payslip never said. The export compares
the two and refuses; that comparison is what is pinned here.

The *loan* posting is covered in KSW_commissions instead, not here: its
account is ``hr.employee.x_loan_acc_no``, declared by KSW_deduction, which
loads after this module — during KSW_payroll's own tests the field does
not exist yet. KSW_commissions loads after both and exercises the same
``ksw.bas.journal.loan_account`` end to end.
"""
import base64
import io
from datetime import date

import openpyxl

from odoo.exceptions import UserError
from odoo.tests.common import TransactionCase


class BasJournalPayrollCommon(TransactionCase):

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        env = cls.env
        cls.rule_basic = cls._rule('BASIC', 'om_hr_payroll.BASIC', 10)
        cls.rule_hra = cls._rule('HRA', 'om_hr_payroll.HRA', 20)
        cls.rule_gross = cls._rule('GROSS', 'om_hr_payroll.GROSS', 30)
        cls.rule_ded = cls._rule('ATTDED', 'om_hr_payroll.DED', 40)
        cls.rule_gosi = cls._rule('GOSI', 'om_hr_payroll.DED', 50)
        cls.rule_net = cls._rule('NET', 'om_hr_payroll.NET', 60)

        # The mapping the accountant would fill in once.
        cls.rule_basic.sudo().write({
            'x_bas_posting': 'expense',
            'x_bas_account_code': '3101010001',
            'x_bas_account_name': 'مصروفات الرواتب الأساسية'})
        cls.rule_hra.sudo().write({
            'x_bas_posting': 'expense',
            'x_bas_account_code': '3101010002',
            'x_bas_account_name': 'مصروفات بدل السكن'})
        cls.rule_ded.sudo().write({
            'x_bas_posting': 'liability',
            'x_bas_account_code': '2107010009',
            'x_bas_account_name': 'خصومات الغياب'})
        cls.rule_gosi.sudo().write({
            'x_bas_posting': 'liability',
            'x_bas_account_code': '2107010007',
            'x_bas_account_name': 'تأمينات اجتماعية مستحقة'})
        cls.rule_net.sudo().write({
            'x_bas_posting': 'payable',
            'x_bas_account_code': '2107010002',
            'x_bas_account_name': 'رواتب مستحقة الدفع'})
        cls.rule_gross.sudo().write({'x_bas_posting': 'none'})

        cls.employee = env['hr.employee'].sudo().create(
            {'name': 'Journal Test Employee'})
        cls.batch = env['hr.payslip.run'].sudo().create({
            'name': 'Journal March 2026',
            'date_start': date(2026, 3, 1),
            'date_end': date(2026, 3, 31),
        })

    @classmethod
    def _rule(cls, code, category_xmlid, sequence):
        rule = cls.env['hr.salary.rule'].sudo().search(
            [('code', '=', code)], limit=1)
        if not rule:
            rule = cls.env['hr.salary.rule'].sudo().create({
                'name': code, 'code': code,
                'category_id': cls.env.ref(category_xmlid).id,
                'sequence': sequence, 'condition_select': 'none',
                'amount_select': 'fix', 'amount_fix': 0,
            })
        return rule

    def _payslip(self, amounts, employee=None):
        """``amounts`` is [(rule, amount), …] in the order they appear."""
        employee = employee or self.employee
        slip = self.env['hr.payslip'].sudo().create({
            'name': 'Slip %s' % employee.name,
            'employee_id': employee.id,
            'date_from': self.batch.date_start,
            'date_to': self.batch.date_end,
            'payslip_run_id': self.batch.id,
            'version_id': employee.current_version_id.id,
        })
        for rule, amount in amounts:
            self.env['hr.payslip.line'].sudo().create({
                'slip_id': slip.id,
                'salary_rule_id': rule.id,
                'code': rule.code, 'name': rule.name,
                'category_id': rule.category_id.id,
                'employee_id': employee.id,
                'version_id': employee.current_version_id.id,
                'amount': amount, 'quantity': 1, 'rate': 100,
                'sequence': rule.sequence,
            })
        return slip

    def _rows(self):
        return self.batch._bas_journal_rows()


class TestPayslipJournalRows(BasJournalPayrollCommon):

    def test_01_earnings_debited_net_credited(self):
        self._payslip([(self.rule_basic, 5000.0), (self.rule_hra, 1250.0),
                       (self.rule_gross, 6250.0), (self.rule_net, 6250.0)])

        rows = self._rows()

        self.assertEqual([r['code'] for r in rows],
                         ['3101010001', '3101010002', '2107010002'])
        self.assertAlmostEqual(rows[0]['debit'], 5000.0)
        self.assertAlmostEqual(rows[1]['debit'], 1250.0)
        self.assertAlmostEqual(rows[2]['credit'], 6250.0)
        # GROSS is a total of the two above it and is never posted.
        self.assertNotIn(6250.0, [r['debit'] for r in rows])

    def test_02_deductions_credited_before_the_net(self):
        self._payslip([(self.rule_basic, 5000.0), (self.rule_gross, 5000.0),
                       (self.rule_ded, -300.0), (self.rule_gosi, -700.0),
                       (self.rule_net, 4000.0)])

        rows = self._rows()

        self.assertEqual([r['code'] for r in rows],
                         ['3101010001', '2107010009', '2107010007',
                          '2107010002'])
        self.assertAlmostEqual(rows[1]['credit'], 300.0)
        self.assertAlmostEqual(rows[2]['credit'], 700.0)
        self.assertAlmostEqual(rows[3]['credit'], 4000.0)  # what he is paid
        self.env['ksw.bas.journal'].check_balanced(rows)

    def test_03_cancelled_payslips_are_not_posted(self):
        self._payslip([(self.rule_basic, 5000.0), (self.rule_gross, 5000.0),
                       (self.rule_net, 5000.0)])
        other = self.env['hr.employee'].sudo().create({'name': 'Cancelled'})
        self._payslip([(self.rule_basic, 9999.0), (self.rule_gross, 9999.0),
                       (self.rule_net, 9999.0)],
                      employee=other).state = 'cancel'

        rows = self._rows()

        self.assertNotIn(9999.0, [r['debit'] for r in rows])
        self.assertAlmostEqual(sum(r['debit'] for r in rows), 5000.0)


class TestPayslipJournalRefusals(BasJournalPayrollCommon):

    def test_11_an_unmapped_rule_refuses(self):
        """The entry would still balance — and pay a net nobody agreed to."""
        bonus = self._rule('JOURNAL_BONUS', 'om_hr_payroll.ALW', 25)
        bonus.sudo().x_bas_posting = 'none'
        self._payslip([(self.rule_basic, 5000.0), (bonus, 500.0),
                       (self.rule_gross, 5500.0), (self.rule_net, 5500.0)])

        with self.assertRaises(UserError) as caught:
            self._rows()
        message = str(caught.exception)
        self.assertIn(self.employee.name, message)
        self.assertIn('JOURNAL_BONUS', message)

    def test_12_no_payable_rule_refuses(self):
        self.rule_net.sudo().x_bas_posting = 'none'
        self._payslip([(self.rule_basic, 5000.0), (self.rule_gross, 5000.0),
                       (self.rule_net, 5000.0)])

        with self.assertRaises(UserError):
            self._rows()


class TestPayslipJournalWorkbook(BasJournalPayrollCommon):

    def test_20_the_wizard_offers_the_mode(self):
        self._payslip([(self.rule_basic, 5000.0), (self.rule_gross, 5000.0),
                       (self.rule_net, 5000.0)])
        wizard = self.env['ksw.bank.file.export.wizard'].sudo().create({
            'payslip_run_id': self.batch.id,
            'export_mode': 'journal_entry',
        })

        action = wizard.action_export()

        attachment = self.env['ir.attachment'].browse(
            int(action['url'].split('/web/content/')[1].split('?')[0]))
        self.assertTrue(attachment.name.startswith('JournalEntry_'))
        sheet = openpyxl.load_workbook(
            io.BytesIO(base64.b64decode(attachment.datas))).active
        self.assertEqual(
            [c.value for c in sheet[2]],
            ['number', 'fdate', 'fcode', 'fname', 'amount', 'tamount',
             'ref', 'remark', 'cost_Center'])
        # Period end, and the voucher number left for BAS to assign.
        self.assertEqual(sheet.cell(3, 2).value.strip(), '31/03/2026')
        self.assertEqual(sheet.cell(3, 1).value.strip(), '')
